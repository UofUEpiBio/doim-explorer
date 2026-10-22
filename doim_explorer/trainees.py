"""Publications for the program's fellows and residents, collected like faculty's.

The roster spreadsheet is internal and is never committed. Each roster member becomes a
directory-shaped record with an affiliation-scoped PubMed author query - the same query
:func:`doim_explorer.directory.build_pubmed_query` gives faculty - and the core collector
runs over them unchanged. The result is its own document pair, so trainee ids never
enter the faculty publication contract: every work it holds is attributed to trainees
through ``faculty_ids`` (the collector's field name) and is tagged ``trainee_ids``.
"""

from __future__ import annotations

import re
import unicodedata
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

from research_explorer.works import build_works_snapshot, split_works_snapshot

from doim_explorer.contracts import ProfileError
from doim_explorer.directory import DEFAULT_PUBMED_AFFILIATIONS

TRAINEE_PUBLICATIONS_DOCUMENT_TYPE = "doim-trainee-publications"
TRAINEE_PUBLICATION_DETAILS_DOCUMENT_TYPE = "doim-trainee-publication-details"
# The roster's column order: Last Name, First Name, Degree, Program, E-Mail Address.
ROSTER_COLUMNS = ("last", "first", "degree", "program", "email")


def norm_name(value: str) -> str:
    """Lower-case ASCII with punctuation as spaces, as the R matcher normalizes names."""

    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z ]+", " ", value.lower())).strip()


def trainee_id(last: str, first: str) -> str:
    """A stable, readable trainee id: ``t-`` plus the normalized last and first names."""

    return "t-" + f"{norm_name(last)} {norm_name(first)}".replace(" ", "-")


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", norm_name(value)).strip("-") or "unassigned"


def read_roster(path: str | Path) -> list[dict[str, str]]:
    """Read the roster's first sheet positionally; the header row is skipped."""

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = workbook.worksheets[0].iter_rows(min_row=2, values_only=True)
    roster = []
    for row in rows:
        values = [str(value).strip() if value is not None else "" for value in row[:5]]
        record = dict(zip(ROSTER_COLUMNS, values + [""] * (5 - len(values))))
        if record["last"] and record["first"]:
            roster.append(record)
    return roster


def build_trainee_pubmed_query(
    last: str,
    first: str,
    email: str = "",
    affiliations: Sequence[str] = DEFAULT_PUBMED_AFFILIATIONS,
) -> str:
    """Faculty's affiliation-scoped PubMed query, on the full first name.

    Initials collide on common names for early-career authors, so trainees are searched
    by full author name: the roster's surname (compound surnames such as "Van Hook" stay
    whole) and the first word of the first name, with or without anything after it -
    a quoted full name is exact, so "Doe, Jane" alone would miss "Doe, Jane A". Full
    author names are indexed from 2002 on, which covers every trainee's career. The
    e-mail local part often carries the name a person publishes under ("Sam.Van.Hook@"
    for roster first name "Reed"); that first name is searched as well.
    """

    def ascii_words(value: str) -> list[str]:
        value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
        return re.findall(r"[A-Za-z'-]+", value)

    surname = " ".join(ascii_words(last))
    firsts = ascii_words(first)[:1]
    local = ascii_words(email.split("@", 1)[0].replace(".", " "))
    if len(local) >= 2 and all(local[0].lower() != name.lower() for name in firsts):
        firsts.append(local[0].capitalize())
    names = [f"{surname}, {name}" for name in firsts] or [surname]
    authors = " OR ".join(
        f'"{name}"[Full Author Name] OR "{name} *"[Full Author Name]' for name in names
    )
    authors = f"({authors})"
    terms = [f'"{term.strip()}"[Affiliation]' for term in affiliations if term.strip()]
    if not terms:
        return authors
    return f'{authors} AND ({" OR ".join(terms)})'


def build_trainee_directory(
    roster: Sequence[Mapping[str, str]],
    directory: Mapping[str, Any],
    manifest: Mapping[str, Any],
) -> dict[str, Any]:
    """A directory-shaped document over the roster: programs are its divisions.

    Collection settings come from the published faculty directory and the affiliation
    terms from the directory manifest, so both groups are searched identically.
    """

    pubmed = manifest.get("pubmed", {})
    pubmed = pubmed if isinstance(pubmed, Mapping) else {}
    affiliations = pubmed.get("affiliations") or DEFAULT_PUBMED_AFFILIATIONS
    # A roster member who is also listed as faculty is one person: compare first and last
    # name words, since faculty names carry middle names and degrees inconsistently.
    faculty_by_key: dict[str, list[str]] = {}
    for member in directory.get("faculty") or []:
        words = norm_name(str(member.get("full_name", "")).split(",", 1)[0]).split()
        if words:
            faculty_by_key.setdefault(f"{words[0]} {words[-1]}", []).append(str(member["id"]))
    programs: dict[str, dict[str, Any]] = {}
    trainees: dict[str, dict[str, Any]] = {}
    for row in roster:
        program = row.get("program") or "Unassigned"
        program_id = _slug(program)
        programs.setdefault(
            program_id,
            {"id": program_id, "name": program, "source_url": "", "faculty_url": ""},
        )
        record_id = trainee_id(row["last"], row["first"])
        if record_id in trainees:
            raise ProfileError(f"the roster lists {record_id} more than once")
        first_words, last_words = norm_name(row["first"]).split(), norm_name(row["last"]).split()
        key = f"{first_words[0]} {last_words[-1]}" if first_words and last_words else ""
        same_person = faculty_by_key.get(key, [])
        trainees[record_id] = {
            "id": record_id,
            "full_name": f"{row['first']} {row['last']}",
            "profile_url": "",
            "division_ids": [program_id],
            "program": program,
            "degree": row.get("degree", ""),
            "faculty_id": same_person[0] if len(same_person) == 1 else None,
            "pubmed_query": build_trainee_pubmed_query(
                row["last"], row["first"], row.get("email", ""), affiliations
            ),
            "collect_publications": True,
        }
    return {
        "settings": dict(directory.get("settings") or {}),
        "affiliations": list(affiliations),
        "divisions": list(programs.values()),
        "faculty": list(trainees.values()),
    }


def build_trainee_publications_snapshot(
    trainee_directory: Mapping[str, Any], **kwargs: Any
) -> dict[str, Any]:
    """Run the faculty collector over trainees and tag every work as theirs.

    Only roster members with at least one publication are named in the result (the
    roster itself stays internal); the rest are counted per program in ``summary``.
    """

    snapshot = build_works_snapshot(dict(trainee_directory), **kwargs)
    snapshot["document_type"] = TRAINEE_PUBLICATIONS_DOCUMENT_TYPE
    for work in snapshot["works"]:
        work["trainee_ids"] = list(work.get("faculty_ids") or [])
        work["program_ids"] = list(work.get("division_ids") or [])
    counts = snapshot.pop("publications_per_faculty")
    snapshot.pop("works_per_researcher", None)
    snapshot["stats"]["trainees_with_publications"] = snapshot["stats"].pop(
        "faculty_with_publications"
    )
    snapshot["stats"].pop("researchers_with_works", None)
    snapshot["health"] = [row for row in snapshot["health"] if row.get("faculty_id") in counts]

    roster = list(trainee_directory["faculty"])
    snapshot["trainees"] = [
        {
            "id": member["id"],
            "name": member["full_name"],
            "degree": member.get("degree", ""),
            "program": member.get("program", ""),
            "program_id": member["division_ids"][0],
            "faculty_id": member.get("faculty_id"),
            "works": counts[member["id"]],
        }
        for member in roster
        if member["id"] in counts
    ]
    by_program: dict[str, dict[str, Any]] = {}
    for member in roster:
        row = by_program.setdefault(
            member.get("program", ""),
            {"program": member.get("program", ""), "roster": 0, "with_publications": 0},
        )
        row["roster"] += 1
        row["with_publications"] += member["id"] in counts
    snapshot["summary"] = {
        "roster": len(roster),
        "with_publications": len(snapshot["trainees"]),
        "by_program": sorted(by_program.values(), key=lambda row: (-row["roster"], row["program"])),
    }
    snapshot["method"] = {
        "query": "full author name (surname, first name, any middle name) at a University "
        "of Utah affiliation, as faculty are searched",
        "affiliations": list(trainee_directory.get("affiliations") or []),
    }
    return snapshot


def split_trainee_publications_snapshot(
    snapshot: Mapping[str, Any],
) -> tuple[dict[str, Any], dict[str, Any]]:
    index, details = split_works_snapshot(dict(snapshot))
    index["document_type"] = TRAINEE_PUBLICATIONS_DOCUMENT_TYPE
    details["document_type"] = TRAINEE_PUBLICATION_DETAILS_DOCUMENT_TYPE
    return index, details
